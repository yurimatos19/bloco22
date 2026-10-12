"""Monta a planilha do Bloco do 22 (xlsx) para subir no Google Drive como Google Sheets."""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter as L

OUT = sys.argv[1]

G900, G800, G600, G100 = '032814', '053B1E', '0B7A35', 'E6F4EA'
Y, Y50, BLUE, BLUE50 = 'FFD21F', 'FFF8D6', '1D3B9A', 'E9EEFB'
PAPER, LINE, MUTED, GREY, WHITE = 'F2F6F0', 'D5DED8', '5B6A61', '9AA59E', 'FFFFFF'

SIZES = ['P', 'M', 'G', 'GG']
STOCK = {'P': 150, 'M': 250, 'G': 400, 'GG': 200}
HEADERS = ['Código', 'Reservado em', 'Nome completo', 'WhatsApp', 'P', 'M', 'G', 'GG', 'Qtd.', 'Total',
           'Status', 'Atendente', 'Retirado em', 'Observações']
WIDTHS = [100, 130, 220, 135, 44, 44, 44, 44, 52, 90, 110, 130, 130, 260]
SIZE_COL = {'P': 'E', 'M': 'F', 'G': 'G', 'GG': 'H'}
EVENTO = 'Quinta-feira, 22/10/2026, às 19h, na Pracinha de Boa Viagem'
BRL = '"R$" #,##0.00'

fill = lambda c: PatternFill('solid', start_color=c, end_color=c)
thin = Side(style='thin', color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
left = Alignment(horizontal='left', vertical='center', wrap_text=True)


def px(ws, col, pixels):
    ws.column_dimensions[L(col)].width = round(pixels / 7, 1)


def h(ws, row, pixels):
    ws.row_dimensions[row].height = round(pixels * 0.75, 1)


def paint(ws, rng, color):
    for row in ws[rng]:
        for c in row:
            c.fill = fill(color)


def merge(ws, rng, value=None, **style):
    ws.merge_cells(rng)
    c = ws[rng.split(':')[0]]
    if value is not None:
        c.value = value
    for k, v in style.items():
        setattr(c, k, v)
    return c


def title(ws, ncols, text, sub):
    last = L(ncols)
    merge(ws, f'A1:{last}1', text, font=Font(name='Montserrat', size=18, bold=True, color=Y), alignment=left)
    merge(ws, f'A2:{last}2', sub, font=Font(name='Roboto', size=10, color=WHITE),
          alignment=Alignment(horizontal='left', vertical='top', wrap_text=True))
    paint(ws, f'A1:{last}2', G800)
    h(ws, 1, 52); h(ws, 2, 28)


def section(ws, rng, text):
    merge(ws, rng, text, font=Font(name='Montserrat', size=12, bold=True, color=G800),
          alignment=Alignment(vertical='bottom'))


def header(ws, rng):
    for row in ws[rng]:
        for c in row:
            c.fill = fill(G800); c.font = Font(name='Roboto', bold=True, color=WHITE); c.alignment = center


wb = Workbook()
normal = NamedStyle(name='base', font=Font(name='Roboto', size=10), alignment=Alignment(vertical='center'))
wb.add_named_style(normal)

painel = wb.active; painel.title = 'Painel'
balcao = wb.create_sheet('Balcão')
reservas = wb.create_sheet('Reservas')
config = wb.create_sheet('Configurações')
guia = wb.create_sheet('Como usar')

# ---------------------------------------------------------------- Configurações
ws = config
ws.sheet_properties.tabColor = BLUE
ws.sheet_view.showGridLines = False
for i, w in enumerate([260, 320, 380], 1): px(ws, i, w)
paint(ws, 'A1:C24', PAPER)
title(ws, 3, 'Configurações', 'Mude os valores nos campos amarelos. O site usa os valores novos na próxima reserva.')
section(ws, 'A4:C4', 'Reservas')
rows = [
    ('Preço por camisa', 22, 'Se mudar o preço, avise quem cuida do site para atualizar o texto.'),
    ('Camisas por pessoa (por WhatsApp)', 1, 'Soma todas as reservas não canceladas do mesmo WhatsApp.'),
    ('Reservas abertas no site', 'Sim', 'Desmarque (ou mude para Não) para encerrar as reservas. O site trava na hora.'),
]
for i, (a, b, c) in enumerate(rows, 5):
    ws[f'A{i}'] = a; ws[f'B{i}'] = b; ws[f'C{i}'] = c
ws['B5'].number_format = BRL
ws['B6'].number_format = '0'
dv = DataValidation(type='list', formula1='"Sim,Não"', allow_blank=False); ws.add_data_validation(dv); dv.add('B7')
section(ws, 'A9:C9', 'Estoque por tamanho')
ws['A10'], ws['B10'] = 'Tamanho', 'Camisas'
header(ws, 'A10:B10')
for i, sz in enumerate(SIZES, 11):
    ws[f'A{i}'] = sz; ws[f'B{i}'] = STOCK[sz]
    ws[f'A{i}'].font = Font(name='Montserrat', bold=True); ws[f'A{i}'].fill = fill(WHITE)
    ws[f'A{i}'].alignment = center; ws[f'B{i}'].alignment = center; ws[f'B{i}'].number_format = '0'
ws['A15'] = 'Total'; ws['B15'] = '=SUM(B11:B14)'
for c in ('A15', 'B15'):
    ws[c].font = Font(name='Roboto', bold=True); ws[c].fill = fill(G100); ws[c].alignment = center
ws['C11'] = 'Para mudar o estoque, edite estes números. O saldo do site se ajusta sozinho.'
section(ws, 'A17:C17', 'Informações para consulta')
info = [('Local de retirada', 'Comitê da Direita Pernambucana'),
        ('Endereço', 'Av. Herculano Bandeira, 766 — Pina, Recife/PE'),
        ('Horário de retirada', 'Segunda a sábado, das 9h às 18h'),
        ('Evento', EVENTO)]
for i, (a, b) in enumerate(info, 18):
    ws[f'A{i}'] = a; ws[f'B{i}'] = b; ws[f'A{i}'].fill = fill(WHITE); ws[f'B{i}'].fill = fill(WHITE)
    ws[f'A{i}'].font = Font(name='Roboto', size=11, bold=True)
for c in ['A5', 'A6', 'A7']:
    ws[c].font = Font(name='Roboto', size=11, bold=True)
for c in ['B5', 'B6', 'B7', 'B11', 'B12', 'B13', 'B14']:
    ws[c].fill = fill(Y50); ws[c].border = BOX; ws[c].font = Font(name='Roboto', size=11, bold=True)
    ws[c].alignment = Alignment(horizontal='left' if c in ('B5', 'B6', 'B7') else 'center', vertical='center')
for c in ['C5', 'C6', 'C7', 'C11']:
    ws[c].font = Font(name='Roboto', size=10, italic=True, color=MUTED); ws[c].alignment = left
for r in range(4, 22): h(ws, r, 30)
for r in (4, 9, 17): h(ws, r, 34)

wb.defined_names['CFG_PRECO'] = DefinedName('CFG_PRECO', attr_text="'Configurações'!$B$5")
wb.defined_names['CFG_MAX'] = DefinedName('CFG_MAX', attr_text="'Configurações'!$B$6")
wb.defined_names['CFG_ABERTAS'] = DefinedName('CFG_ABERTAS', attr_text="'Configurações'!$B$7")
wb.defined_names['CFG_ESTOQUE'] = DefinedName('CFG_ESTOQUE', attr_text="'Configurações'!$A$11:$B$14")

# ---------------------------------------------------------------- Reservas
ws = reservas
ws.sheet_properties.tabColor = G600
for i, w in enumerate(WIDTHS, 1): px(ws, i, w)
for i, t in enumerate(HEADERS, 1): ws.cell(1, i, t)
header(ws, 'A1:N1')
for col in 'EFGH': ws[f'{col}1'].fill = fill(G600)
h(ws, 1, 36)
ws.freeze_panes = 'B2'
ws.auto_filter.ref = 'A1:N2000'
dv = DataValidation(type='list', formula1='"Reservado,Entregue,Cancelado"', allow_blank=True)
ws.add_data_validation(dv); dv.add('K2:K2000')
ws.conditional_formatting.add('A2:N2000', FormulaRule(formula=['$K2="Cancelado"'], font=Font(color=GREY, strike=True)))
ws.conditional_formatting.add('K2:K2000', FormulaRule(formula=['$K2="Entregue"'], fill=fill(G100), font=Font(color=G800, bold=True)))
ws.conditional_formatting.add('K2:K2000', FormulaRule(formula=['$K2="Reservado"'], fill=fill(Y50), font=Font(color='5C4A00', bold=True)))
ws.conditional_formatting.add('A2:N2000', FormulaRule(formula=['AND($A2<>"",MOD(ROW(),2)=1)'], fill=fill(PAPER)))

# ---------------------------------------------------------------- Painel
ws = painel
ws.sheet_properties.tabColor = Y
ws.sheet_view.showGridLines = False
for i in range(1, 9): px(ws, i, 128)
paint(ws, 'A1:H24', PAPER)
title(ws, 8, 'Bloco do 22 — Painel de reservas', f'Atualiza sozinho a cada reserva. Evento: {EVENTO}.')
R, S = 'Reservas!', 'Reservas!K2:K'
ST = 'SUM(INDEX(CFG_ESTOQUE,0,2))'
cards = [
    (4, 1, 'Reservas ativas', f'=COUNTIFS({R}A2:A,"<>",{S},"<>Cancelado")', '0', 'pedidos não cancelados'),
    (4, 3, 'Camisas reservadas', f'=SUMIFS({R}I2:I,{S},"<>Cancelado")', '0', f'="de "&{ST}&" camisas"'),
    (4, 5, 'Entregues', f'=SUMIFS({R}I2:I,{S},"Entregue")', '0', 'camisas já retiradas'),
    (4, 7, 'A entregar', '=C5-E5', '0', 'reservadas e ainda não retiradas'),
    (8, 1, 'Disponíveis no site', f'=MAX(0,{ST}-C5)', '0', 'camisas que ainda dá para reservar'),
    (8, 3, 'Valor a receber', f'=SUMIFS({R}J2:J,{S},"Reservado")', BRL, 'das reservas pendentes'),
    (8, 5, 'Valor recebido', f'=SUMIFS({R}J2:J,{S},"Entregue")', BRL, 'das camisas entregues'),
    (8, 7, 'Canceladas', f'=COUNTIF({S},"Cancelado")', '0', 'reservas canceladas'),
]
for row, col, label, formula, fmt, note in cards:
    a, b = L(col), L(col + 1)
    for r in range(row, row + 3):
        for c in (a, b):
            ws[f'{c}{r}'].fill = fill(WHITE)
            ws[f'{c}{r}'].border = Border(left=thin if c == a else None, right=thin if c == b else None,
                                          top=thin if r == row else None, bottom=thin if r == row + 2 else None)
    merge(ws, f'{a}{row}:{b}{row}', label, font=Font(name='Roboto', size=10, bold=True, color=MUTED), alignment=left)
    v = merge(ws, f'{a}{row+1}:{b}{row+1}', formula,
              font=Font(name='Montserrat', size=24, bold=True, color=BLUE if label == 'A entregar' else G800), alignment=left)
    v.number_format = fmt
    merge(ws, f'{a}{row+2}:{b}{row+2}', note, font=Font(name='Roboto', size=9, color=MUTED), alignment=left)
h(ws, 3, 16); h(ws, 7, 14); h(ws, 11, 14)
for r in (4, 8): h(ws, r, 26); h(ws, r + 1, 46); h(ws, r + 2, 22)
section(ws, 'A12:H12', 'Estoque por tamanho'); h(ws, 12, 34)
for i, t in enumerate(['Tamanho', 'Estoque', 'Reservadas', 'Entregues', 'A entregar', 'Disponíveis', 'Ocupação', '%'], 1):
    ws.cell(13, i, t)
header(ws, 'A13:H13'); h(ws, 13, 30)
for i, sz in enumerate(SIZES):
    r = 14 + i; c = SIZE_COL[sz]
    vals = [sz, f'=INDEX(CFG_ESTOQUE,{i+1},2)', f'=SUMIFS({R}{c}2:{c},{S},"<>Cancelado")',
            f'=SUMIFS({R}{c}2:{c},{S},"Entregue")', f'=C{r}-D{r}', f'=MAX(0,B{r}-C{r})',
            f'=SPARKLINE(C{r},{{"charttype","bar";"max",MAX(1,B{r});"color1","#{G600}"}})', f'=IFERROR(C{r}/B{r},0)']
    for j, v in enumerate(vals, 1): ws.cell(r, j, v)
tot = 14 + len(SIZES)
for j, v in enumerate(['Total', '=SUM(B14:B17)', '=SUM(C14:C17)', '=SUM(D14:D17)', '=SUM(E14:E17)', '=SUM(F14:F17)',
                       f'=SPARKLINE(C{tot},{{"charttype","bar";"max",MAX(1,B{tot});"color1","#{G800}"}})',
                       f'=IFERROR(C{tot}/B{tot},0)'], 1):
    ws.cell(tot, j, v)
for r in range(14, tot + 1):
    h(ws, r, 30)
    for j in range(1, 9):
        c = ws.cell(r, j)
        c.fill = fill(G100 if r == tot else WHITE); c.alignment = center; c.border = Border(bottom=thin, top=thin)
        c.font = Font(name='Montserrat' if j == 1 else 'Roboto', size=11, bold=(j in (1, 6) or r == tot),
                      color=G800 if j == 6 else '0F1D15')
        c.number_format = '0%' if j == 8 else '0'
ws.conditional_formatting.add('F14:F17', FormulaRule(formula=['AND($B14>0,$F14<=20)'], fill=fill(BLUE50), font=Font(color=BLUE, bold=True)))
merge(ws, f'A{tot+2}:H{tot+2}',
      'Para dar baixa numa retirada, use a aba Balcão. Para cancelar uma reserva, mude o Status na aba Reservas.',
      font=Font(name='Roboto', size=10, color=MUTED), alignment=left)

# ---------------------------------------------------------------- Balcão
ws = balcao
ws.sheet_properties.tabColor = G800
ws.sheet_view.showGridLines = False
for i, w in enumerate(WIDTHS, 1): px(ws, i, 160 if i == 1 else w)
title(ws, 14, 'Balcão de retirada',
      'Digite o código ou o WhatsApp, confira a reserva, receba o pagamento e marque Entregue.')
for i, t in enumerate(['Código ou WhatsApp', 'Atendente', 'Entregue'], 3):
    ws[f'A{i}'] = t; ws[f'A{i}'].font = Font(name='Roboto', size=11, bold=True)
    ws[f'A{i}'].alignment = Alignment(horizontal='right', vertical='center')
for rng, size in (('B3:C3', 14), ('B4:C4', 11)):
    ws.merge_cells(rng)
    a, b = rng.split(':')
    for c in (a, b):
        ws[c].fill = fill(Y50)
        ws[c].border = Border(top=Side('medium', color=G600), bottom=Side('medium', color=G600),
                              left=Side('medium', color=G600) if c == a else None,
                              right=Side('medium', color=G600) if c == b else None)
    ws[a].font = Font(name='Roboto Mono' if size == 14 else 'Roboto', size=size, bold=size == 14)
ws['B5'] = 'Não'
ws['B5'].fill = fill(Y50); ws['B5'].border = BOX; ws['B5'].font = Font(name='Roboto', size=11, bold=True)
dv = DataValidation(type='list', formula1='"Não,Sim"', allow_blank=False); ws.add_data_validation(dv); dv.add('B5')
ws['D3'] = 'Digite e aperte Enter'
ws['D5'] = 'Marque depois de receber o pagamento e entregar todas as camisas'
for c in ('D3', 'D5'): ws[c].font = Font(name='Roboto', italic=True, color=MUTED)
merge(ws, 'B6:N6', None, font=Font(name='Roboto', size=11, bold=True), alignment=left)
for r in (3, 4, 5, 6): h(ws, r, 34)
for i, t in enumerate(HEADERS, 1): ws.cell(8, i, t)
header(ws, 'A8:N8'); h(ws, 8, 32)
ws['A9'] = ('=ARRAYFORMULA(IF($B$3="","",IFERROR(FILTER(Reservas!A2:N,'
            '(Reservas!A2:A=UPPER(TRIM($B$3)))+'
            'IF(LEN(REGEXREPLACE($B$3&"","\\D",""))>=8,RIGHT(REGEXREPLACE(Reservas!D2:D&"","\\D",""),'
            'LEN(REGEXREPLACE($B$3&"","\\D","")))=REGEXREPLACE($B$3&"","\\D",""),FALSE)'
            '),"Nenhuma reserva encontrada para esse código ou WhatsApp.")))')
for r in range(9, 29):
    for j in range(1, 15):
        c = ws.cell(r, j); c.fill = fill(WHITE); c.border = Border(bottom=thin)
        c.alignment = Alignment(horizontal='center' if 5 <= j <= 11 else 'left', vertical='center')
    ws.cell(r, 1).font = Font(name='Roboto Mono', bold=True)
    ws.cell(r, 2).number_format = 'dd/mm/yyyy hh:mm'
    ws.cell(r, 13).number_format = 'dd/mm/yyyy hh:mm'
    ws.cell(r, 10).number_format = BRL
ws.conditional_formatting.add('K9:K28', FormulaRule(formula=['$K9="Entregue"'], fill=fill(G100), font=Font(color=G800, bold=True)))
ws.conditional_formatting.add('K9:K28', FormulaRule(formula=['$K9="Reservado"'], fill=fill(Y50), font=Font(color='5C4A00', bold=True)))
ws.freeze_panes = 'A9'

# ---------------------------------------------------------------- Como usar
ws = guia
ws.sheet_properties.tabColor = MUTED
ws.sheet_view.showGridLines = False
px(ws, 1, 56); px(ws, 2, 760)
paint(ws, 'A1:B30', PAPER)
title(ws, 2, 'Como usar', 'Guia rápido para a equipe do comitê.')
blocks = [
    ('Na retirada (aba Balcão)', [
        'Peça o código da reserva (começa com B22-) ou o WhatsApp da pessoa e digite no campo amarelo.',
        'Confira nome, tamanhos e total. Se o Status for Entregue, a camisa já foi retirada: não entregue de novo.',
        'Receba o pagamento e entregue TODAS as camisas do pedido.',
        'Escreva seu nome em Atendente e marque Entregue. A planilha registra a data e a hora.',
    ]),
    ('Se alguém desistir', [
        'Na aba Reservas, mude o Status da reserva para Cancelado. As camisas voltam para o estoque do site.',
    ]),
    ('Estoque, preço e encerramento (aba Configurações)', [
        'Para mudar o estoque de um tamanho, edite o número na tabela Estoque por tamanho.',
        'Para encerrar as reservas, desmarque Reservas abertas no site. Para reabrir, marque de novo.',
    ]),
    ('Acompanhamento (aba Painel)', [
        'Mostra reservas, camisas entregues, valores e quanto resta de cada tamanho. Tudo se atualiza sozinho.',
    ]),
]
r = 4
for t, steps in blocks:
    section(ws, f'A{r}:B{r}', t); h(ws, r, 34); r += 1
    for i, s in enumerate(steps, 1):
        ws[f'A{r}'] = i; ws[f'B{r}'] = s
        ws[f'A{r}'].fill = fill(Y); ws[f'A{r}'].font = Font(name='Montserrat', bold=True, color=G900); ws[f'A{r}'].alignment = center
        ws[f'B{r}'].fill = fill(WHITE); ws[f'B{r}'].font = Font(name='Roboto', size=11); ws[f'B{r}'].alignment = left
        h(ws, r, 32); r += 1
    r += 1

wb.active = 0
wb.save(OUT)
print('ok', OUT)
